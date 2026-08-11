"""Generate deterministic Azure SQL DDL from the supplied SAP DDIC workbook."""

from __future__ import annotations

import argparse
import hashlib
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = (
    ROOT
    / "specs"
    / "015-sap-hana-private-connectivity"
    / "SAP_Force_Equipment_Table_Relationship_Complete.xlsx"
)
DEFAULT_OUTPUT = Path(__file__).resolve().parent / "migrations" / "001_sap_schema.sql"

SOURCE_SHEETS = {
    "/ISDFPS/FORCE": " ( ISDFPS FORCE)",
    "HRP1000": "HRP1000",
    "HRP1001": "HRP1001",
    "EQUI": "EQUI",
    "EQUZ": "EQUZ",
    "EQKT": "EQKT",
    "MARA": "MARA",
    "ILOA": "ILOA",
    "IFLOT": "IFLOT",
    "JEST": "JEST",
    "TJ02T": "TJO2T",
    "TJ30T": " TJ30T",
}

SQL_TABLE_NAMES = {
    "/ISDFPS/FORCE": "ISDFPS_FORCE",
}

HEADER_ALIASES = {
    "field": {"field", "field name"},
    "key": {"key"},
    "data_element": {"data element"},
    "data_type": {"data type"},
    "length": {"length"},
    "decimals": {"decimals"},
    "description": {"short description"},
}

CHARACTER_TYPES = frozenset({"CHAR", "UNIT"})
FIXED_CHARACTER_TYPES = frozenset({"CLNT", "CUKY", "DATS", "LANG", "NUMC", "TIMS"})
DECIMAL_TYPES = frozenset({"CURR", "DEC", "QUAN"})
INTEGER_TYPES = {"INT1": "tinyint", "INT2": "smallint", "INT8": "bigint"}
SUPPORTED_TYPES = (
    CHARACTER_TYPES
    | FIXED_CHARACTER_TYPES
    | DECIMAL_TYPES
    | frozenset(INTEGER_TYPES)
    | {"RAW"}
)


class WorkbookContractError(ValueError):
    """Raised when the workbook cannot be converted without guessing."""


@dataclass(frozen=True)
class SapField:
    ordinal: int
    name: str
    is_key: bool
    data_element: str
    data_type: str
    length: int
    decimals: int
    description: str


@dataclass(frozen=True)
class SapTable:
    source_name: str
    sql_name: str
    fields: tuple[SapField, ...]

    @property
    def key_fields(self) -> tuple[SapField, ...]:
        return tuple(field for field in self.fields if field.is_key)


def _normalized_header(value: Any) -> str:
    return " ".join(str(value or "").strip().lower().split())


def _find_header(worksheet: Any) -> tuple[int, dict[str, int]]:
    for row in worksheet.iter_rows():
        values = [_normalized_header(cell.value) for cell in row]
        columns: dict[str, int] = {}
        for canonical, aliases in HEADER_ALIASES.items():
            for index, value in enumerate(values, start=1):
                if value in aliases:
                    columns[canonical] = index
                    break
        if set(columns) == set(HEADER_ALIASES):
            return row[0].row, columns
    raise WorkbookContractError(f"Could not find a DDIC header in sheet {worksheet.title!r}.")


def _integer(value: Any, *, context: str) -> int:
    try:
        parsed = int(value)
    except (TypeError, ValueError) as error:
        raise WorkbookContractError(f"{context} must be an integer; got {value!r}.") from error
    if parsed < 0:
        raise WorkbookContractError(f"{context} must not be negative.")
    return parsed


def _parse_table(source_name: str, worksheet: Any) -> SapTable:
    header_row, columns = _find_header(worksheet)
    fields: list[SapField] = []
    seen: set[str] = set()

    for row_number in range(header_row + 1, worksheet.max_row + 1):
        raw_name = worksheet.cell(row_number, columns["field"]).value
        raw_type = worksheet.cell(row_number, columns["data_type"]).value
        if raw_name in (None, "") and raw_type in (None, ""):
            continue

        name = str(raw_name or "").strip()
        data_type = str(raw_type or "").strip().upper()
        if name.startswith(".") or data_type == "STRU":
            continue
        if not name or not data_type:
            raise WorkbookContractError(
                f"Incomplete field definition in {source_name} at workbook row {row_number}."
            )
        if data_type not in SUPPORTED_TYPES:
            raise WorkbookContractError(
                f"Unsupported SAP type {data_type!r} for {source_name}.{name}."
            )
        normalized_name = name.upper()
        if normalized_name in seen:
            raise WorkbookContractError(f"Duplicate field {source_name}.{name}.")
        seen.add(normalized_name)

        key_marker = worksheet.cell(row_number, columns["key"]).value
        fields.append(
            SapField(
                ordinal=len(fields) + 1,
                name=name,
                is_key=str(key_marker or "").strip().upper() in {"X", "✔"},
                data_element=str(
                    worksheet.cell(row_number, columns["data_element"]).value or ""
                ).strip(),
                data_type=data_type,
                length=_integer(
                    worksheet.cell(row_number, columns["length"]).value,
                    context=f"Length for {source_name}.{name}",
                ),
                decimals=_integer(
                    worksheet.cell(row_number, columns["decimals"]).value,
                    context=f"Decimals for {source_name}.{name}",
                ),
                description=str(
                    worksheet.cell(row_number, columns["description"]).value or ""
                ).strip(),
            )
        )

    if not fields:
        raise WorkbookContractError(f"No physical fields found for {source_name}.")
    if not any(field.is_key for field in fields):
        raise WorkbookContractError(f"No key fields found for {source_name}.")

    return SapTable(
        source_name=source_name,
        sql_name=SQL_TABLE_NAMES.get(source_name, source_name),
        fields=tuple(fields),
    )


def load_contract(workbook_path: Path = DEFAULT_WORKBOOK) -> tuple[SapTable, ...]:
    workbook = load_workbook(workbook_path, read_only=False, data_only=False)
    try:
        missing = [sheet for sheet in SOURCE_SHEETS.values() if sheet not in workbook.sheetnames]
        if missing:
            raise WorkbookContractError(f"Workbook is missing expected sheets: {missing!r}.")
        formulas = [
            f"{worksheet.title}!{cell.coordinate}"
            for worksheet in workbook.worksheets
            for row in worksheet.iter_rows()
            for cell in row
            if cell.data_type == "f"
        ]
        if formulas:
            raise WorkbookContractError(
                "DDIC generation does not accept formula-dependent values: " + ", ".join(formulas)
            )
        return tuple(
            _parse_table(source_name, workbook[sheet_name])
            for source_name, sheet_name in SOURCE_SHEETS.items()
        )
    finally:
        workbook.close()


def _quote_identifier(value: str) -> str:
    return f"[{value.replace(']', ']]')}]"


def _sql_string(value: str) -> str:
    return "N'" + value.replace("'", "''") + "'"


def _sql_type(field: SapField) -> str:
    if field.length <= 0:
        raise WorkbookContractError(f"{field.name} has invalid physical length {field.length}.")
    if field.data_type in CHARACTER_TYPES:
        return f"nvarchar({field.length})"
    if field.data_type in FIXED_CHARACTER_TYPES:
        return f"char({field.length})"
    if field.data_type in DECIMAL_TYPES:
        if not 1 <= field.length <= 38 or field.decimals > field.length:
            raise WorkbookContractError(
                f"{field.name} has unsupported decimal shape ({field.length}, {field.decimals})."
            )
        return f"decimal({field.length}, {field.decimals})"
    if field.data_type == "RAW":
        return f"varbinary({field.length})"
    return INTEGER_TYPES[field.data_type]


def _render_table(table: SapTable) -> str:
    qualified_name = f"[sap].{_quote_identifier(table.sql_name)}"
    definitions = [
        f"        {_quote_identifier(field.name)} {_sql_type(field)} "
        f"{'NOT NULL' if field.is_key else 'NULL'}"
        for field in table.fields
    ]
    key_columns = ", ".join(_quote_identifier(field.name) for field in table.key_fields)
    definitions.append(
        f"        CONSTRAINT {_quote_identifier('PK_' + table.sql_name)} "
        f"PRIMARY KEY ({key_columns})"
    )
    body = ",\n".join(definitions)
    return (
        f"IF OBJECT_ID(N'{qualified_name}', N'U') IS NULL\n"
        "BEGIN\n"
        f"    CREATE TABLE {qualified_name} (\n{body}\n    );\n"
        "END;\nGO"
    )


def _render_metadata(tables: tuple[SapTable, ...], workbook_hash: str) -> str:
    rows = []
    for table in tables:
        for field in table.fields:
            values = (
                _sql_string(table.source_name),
                _sql_string(table.sql_name),
                str(field.ordinal),
                _sql_string(field.name),
                "1" if field.is_key else "0",
                _sql_string(field.data_element),
                _sql_string(field.data_type),
                str(field.length),
                str(field.decimals),
                _sql_string(field.description),
                f"'{workbook_hash}'",
            )
            rows.append("        (" + ", ".join(values) + ")")
    joined_rows = ",\n".join(rows)
    return f"""IF OBJECT_ID(N'[emulator].[DDIC_FIELD]', N'U') IS NULL
BEGIN
    CREATE TABLE [emulator].[DDIC_FIELD] (
        [SOURCE_TABLE] nvarchar(128) NOT NULL,
        [SQL_TABLE] sysname NOT NULL,
        [ORDINAL] int NOT NULL,
        [FIELD_NAME] sysname NOT NULL,
        [IS_KEY] bit NOT NULL,
        [DATA_ELEMENT] nvarchar(128) NOT NULL,
        [SAP_DATA_TYPE] varchar(16) NOT NULL,
        [LENGTH] int NOT NULL,
        [DECIMALS] int NOT NULL,
        [DESCRIPTION] nvarchar(1000) NOT NULL,
        [SOURCE_WORKBOOK_SHA256] char(64) NOT NULL,
        CONSTRAINT [PK_DDIC_FIELD]
            PRIMARY KEY ([SOURCE_WORKBOOK_SHA256], [SOURCE_TABLE], [ORDINAL])
    );
END;
GO

IF NOT EXISTS (
    SELECT 1
    FROM [emulator].[DDIC_FIELD]
    WHERE [SOURCE_WORKBOOK_SHA256] = '{workbook_hash}'
)
BEGIN
    INSERT INTO [emulator].[DDIC_FIELD] (
        [SOURCE_TABLE], [SQL_TABLE], [ORDINAL], [FIELD_NAME], [IS_KEY],
        [DATA_ELEMENT], [SAP_DATA_TYPE], [LENGTH], [DECIMALS], [DESCRIPTION],
        [SOURCE_WORKBOOK_SHA256]
    ) VALUES
{joined_rows};
END;
GO"""


def render_schema(workbook_path: Path = DEFAULT_WORKBOOK) -> str:
    tables = load_contract(workbook_path)
    workbook_hash = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    table_sql = "\n\n".join(_render_table(table) for table in tables)
    metadata_sql = _render_metadata(tables, workbook_hash)
    return f"""-- Generated by database_emulator/generate_schema.py. Do not edit manually.
-- Source workbook SHA-256: {workbook_hash}
SET ANSI_NULLS ON;
SET QUOTED_IDENTIFIER ON;
GO

IF NOT EXISTS (SELECT 1 FROM sys.schemas WHERE name = N'sap')
    EXEC(N'CREATE SCHEMA [sap] AUTHORIZATION [dbo]');
GO
IF NOT EXISTS (SELECT 1 FROM sys.schemas WHERE name = N'emulator')
    EXEC(N'CREATE SCHEMA [emulator] AUTHORIZATION [dbo]');
GO

{table_sql}

{metadata_sql}
"""


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()

    rendered = render_schema(args.workbook)
    if args.check:
        if not args.output.exists() or args.output.read_text(encoding="utf-8") != rendered:
            parser.error(f"{args.output} is missing or does not match the workbook")
        return 0

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(rendered, encoding="utf-8", newline="\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())